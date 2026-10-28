"""Build a deck from the example slides of a corporate .pptx: copy, reorder, fill, drop the rest.

Usage:
  python3 pptx_fill.py SOURCE.pptx PLAN.json OUTPUT.pptx

PLAN.json (shape ids and paragraph indexes come from pptx_inventory.py on SOURCE):
{
  "slides": [
    { "from": 1, "texts": { "8": "Annual Report 2026", "18": "Berlin, 28.10.2026" },
      "notes": "Welcome" },
    { "from": 2, "texts": { "3": ["Market", "Products", "Team", "Outlook"] } },
    { "from": 5, "texts": { "4": { "0": "Our strengths" } },          # only paragraph 0
      "replace": { "01 TITEL": "01 COMPANY" },                        # substring, any shape
      "tables": { "9": [["Year", "Revenue"], ["2025", "1.2 M"]] },    # rows x cells, null keeps
      "images": { "12": "/mnt/data/photo.png" },                       # swap picture, keeps frame
      "delete": [6] },                                                 # remove dummy shapes
    { "from": 5, "texts": { "4": "Second slide in the same design" } }
  ]
}

- Output = exactly the listed slides, in that order. A source slide may be used several times (it is
  copied, images/charts included); source slides that are not listed are removed.
- Text values: a string ("\\n" = new paragraph, "<br>" = line break inside a paragraph), a list of
  paragraphs (null = keep that paragraph), or {"<para index>": text} to change single paragraphs.
  Every paragraph keeps the formatting of the paragraph it replaces; extra paragraphs copy the last one.
- Shape ids are those of the SOURCE slide; they stay the same on every copy.
Prints a short report and then the leftover-sample-text check (exit 2 if sample text is left).
"""
import copy
import json
import os
import re
import subprocess
import sys

import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SHARED_RELS = {RT.IMAGE, RT.MEDIA, RT.VIDEO, RT.AUDIO, RT.SLIDE_LAYOUT, RT.HYPERLINK, RT.SLIDE,
               RT.NOTES_MASTER, RT.SLIDE_MASTER, RT.THEME}
SKIP_RELS = {RT.NOTES_SLIDE}
PPTX_V1 = int(pptx.__version__.split(".")[0]) >= 1


# ---------------------------------------------------------------- slide copy

def _remap_rids(element, mapping):
    for el in element.iter():
        for attr, val in list(el.attrib.items()):
            if attr.startswith("{%s}" % R_NS) and val in mapping:
                el.set(attr, mapping[val])


def _clone_part(part, package):
    """Deep-copy a non-shared part (chart, SmartArt data, embedded workbook ...) with its own rels."""
    template = re.sub(r"\d*(\.\w+)$", r"%d\1", str(part.partname))
    partname = package.next_partname(template)
    if PPTX_V1:
        new = type(part).load(partname, part.content_type, package, part.blob)
    else:                                    # python-pptx < 1.0 has blob before package
        new = type(part).load(partname, part.content_type, part.blob, package)
    mapping = {}
    for rid, rel in part.rels.items():
        if rel.is_external:
            mapping[rid] = new.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            target = rel.target_part
            if rel.reltype not in SHARED_RELS:
                target = _clone_part(target, package)
            mapping[rid] = new.rels.get_or_add(rel.reltype, target)
    if hasattr(new, "_element"):
        _remap_rids(new._element, mapping)
    elif mapping and any(k != v for k, v in mapping.items()):
        print(f"  warning: could not remap relationships inside {partname}", file=sys.stderr)
    return new


def duplicate_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    package = prs.part.package
    if sum(1 for p in package.iter_parts() if p.partname == new.part.partname) > 1:
        new.part.partname = package.next_partname("/ppt/slides/slide%d.xml")  # deck with gaps in numbering
    root, src_root = new._element, src._element
    # Swap the content but keep the slide's own spTree element: python-pptx caches it.
    for k, v in src_root.attrib.items():
        root.set(k, v)
    tree, src_tree = root.cSld.spTree, src_root.cSld.spTree
    for child in list(tree):
        tree.remove(child)
    for child in src_tree:
        tree.append(copy.deepcopy(child))
    for k, v in src_root.cSld.attrib.items():
        root.cSld.set(k, v)
    for child in list(root.cSld):
        if child is not tree:
            root.cSld.remove(child)
    for child in src_root.cSld:
        if child is src_tree:
            continue
        clone = copy.deepcopy(child)
        if src_root.cSld.index(child) < src_root.cSld.index(src_tree):
            tree.addprevious(clone)          # bg
        else:
            root.cSld.append(clone)          # custDataLst, controls, extLst
    for child in list(root):
        if child is not root.cSld:
            root.remove(child)
    for child in src_root:
        if child is not src_root.cSld:
            root.append(copy.deepcopy(child))  # clrMapOvr, transition, timing, extLst
    mapping = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in SKIP_RELS or rel.reltype == RT.SLIDE_LAYOUT:
            continue
        if rel.is_external:
            mapping[rid] = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        elif rel.reltype in SHARED_RELS:
            mapping[rid] = new.part.rels.get_or_add(rel.reltype, rel.target_part)
        else:
            mapping[rid] = new.part.rels.get_or_add(rel.reltype, _clone_part(rel.target_part, prs.part.package))
    _remap_rids(root, mapping)
    if src.has_notes_slide:
        new.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    return new


def set_order(prs, slides):
    """Keep exactly `slides` in this order; drop every other slide."""
    sld_ids = prs.slides._sldIdLst
    by_part = {}
    for sld_id in list(sld_ids):
        by_part[prs.part.related_part(sld_id.rId)] = sld_id
        sld_ids.remove(sld_id)
    keep = set()
    for slide in slides:
        sld_ids.append(by_part[slide.part])
        keep.add(slide.part)
    for part, sld_id in by_part.items():
        if part not in keep:
            prs.part.drop_rel(sld_id.rId)


# ---------------------------------------------------------------- text

def _run_props(p):
    r = p._p.find(qn("a:r"))
    if r is not None and r.find(qn("a:rPr")) is not None:
        return copy.deepcopy(r.find(qn("a:rPr")))
    end = p._p.find(qn("a:endParaRPr"))
    if end is not None:
        rpr = copy.deepcopy(end)
        rpr.tag = qn("a:rPr")
        return rpr
    return None


def set_paragraph(p, text):
    """Replace the text of one paragraph, keeping paragraph and first-run formatting."""
    rpr = _run_props(p)
    pel = p._p
    for child in list(pel):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            pel.remove(child)
    end = pel.find(qn("a:endParaRPr"))
    for i, part in enumerate(re.split(r"<br\s*/?>|\v", str(text))):
        if i:
            br = pel.makeelement(qn("a:br"), {})
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            _insert(pel, br, end)
        r = pel.makeelement(qn("a:r"), {})
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = r.makeelement(qn("a:t"), {})
        t.text = part
        r.append(t)
        _insert(pel, r, end)


def _insert(parent, el, before):
    if before is not None:
        before.addprevious(el)
    else:
        parent.append(el)


def fill_text_frame(tf, value):
    paras = list(tf.paragraphs)
    if isinstance(value, dict):
        for idx, text in value.items():
            i = int(idx)
            if i >= len(paras):
                raise ValueError(f"paragraph {i} does not exist (shape has {len(paras)})")
            set_paragraph(paras[i], text)
        return
    wanted = value if isinstance(value, list) else str(value).split("\n")
    # Trailing empty paragraphs of the sample are dropped unless asked for.
    while len(paras) > 1 and len(paras) > len(wanted) and not paras[-1].text.strip():
        tf._txBody.remove(paras[-1]._p)
        paras.pop()
    for i, text in enumerate(wanted):
        if i < len(paras):
            if text is not None:
                set_paragraph(paras[i], text)
            continue
        base = paras[-1]._p
        clone = copy.deepcopy(base)
        base.addnext(clone)
        paras = list(tf.paragraphs)
        set_paragraph(paras[i], "" if text is None else text)
    for p in paras[len(wanted):]:
        tf._txBody.remove(p._p)


def replace_in_paragraph(p, old, new):
    full = "".join(r.text for r in p.runs)
    if old not in full:
        return 0
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return 1
    set_paragraph(p, full.replace(old, new))
    return 1


def all_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from all_shapes(sh.shapes)


def text_frames(shape):
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        yield shape.text_frame
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame


# ---------------------------------------------------------------- per slide

def apply(slide, spec, report):
    shapes = {str(sh.shape_id): sh for sh in all_shapes(slide.shapes)}

    def shape(sid):
        if str(sid) not in shapes:
            raise ValueError(f"shape id {sid} not found (ids: {', '.join(shapes)})")
        return shapes[str(sid)]

    for sid, value in spec.get("texts", {}).items():
        sh = shape(sid)
        if not getattr(sh, "has_text_frame", False) or not sh.has_text_frame:
            raise ValueError(f"shape id {sid} has no text (use 'tables' or 'images')")
        fill_text_frame(sh.text_frame, value)
    for sid, rows in spec.get("tables", {}).items():
        sh = shape(sid)
        if not getattr(sh, "has_table", False) or not sh.has_table:
            raise ValueError(f"shape id {sid} is not a table")
        table = sh.table
        tr_list = list(table._tbl.tr_lst)
        while len(tr_list) < len(rows):          # grow: copy the last row
            clone = copy.deepcopy(tr_list[-1])
            tr_list[-1].addnext(clone)
            tr_list = list(table._tbl.tr_lst)
        for ri, row in enumerate(rows):
            for ci, value in enumerate(row):
                if value is not None and ci < len(table.columns):
                    fill_text_frame(table.cell(ri, ci).text_frame, value)
        if spec.get("trim_tables", True):
            for tr in tr_list[len(rows):]:
                table._tbl.remove(tr)
    for old, new in spec.get("replace", {}).items():
        hits = sum(replace_in_paragraph(p, old, new)
                   for sh in all_shapes(slide.shapes) for tf in text_frames(sh) for p in tf.paragraphs)
        if not hits:
            report.append(f"  warning: '{old}' not found")
    for sid, path in spec.get("images", {}).items():
        sh = shape(sid)
        if sh.shape_type != MSO_SHAPE_TYPE.PICTURE and not sh._element.xpath(".//a:blip"):
            raise ValueError(f"shape id {sid} is not a picture; delete it and keep the frame, or pick a picture")
        _, rid = slide.part.get_or_add_image_part(path)
        blip = sh._element.xpath(".//a:blip")[0]
        old_rid = blip.get(qn("r:embed"))
        blip.set(qn("r:embed"), rid)
        if old_rid and old_rid != rid:
            slide.part.drop_rel(old_rid)     # only drops it when nothing else uses the old image
    for sid in spec.get("delete", []):
        el = shape(sid)._element
        el.getparent().remove(el)
        for rid in {v for e in el.iter() for k, v in e.attrib.items() if k.startswith("{%s}" % R_NS)}:
            slide.part.drop_rel(rid)         # only drops it when nothing else on the slide uses it
    if "notes" in spec:
        slide.notes_slide.notes_text_frame.text = spec["notes"] or ""


def main():
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    src_path, plan_path, out_path = sys.argv[1:]
    if os.path.abspath(src_path) == os.path.abspath(out_path):
        sys.exit("Refusing to overwrite the source; pick a new output name.")
    try:
        with open(plan_path, encoding="utf-8") as fh:
            plan = json.load(fh)
    except json.JSONDecodeError as exc:
        sys.exit(f"{plan_path} is not valid JSON: {exc}")
    prs = Presentation(src_path)
    originals = list(prs.slides)
    used, result = set(), []
    for n, spec in enumerate(plan["slides"], 1):
        src = int(spec["from"])
        if not 1 <= src <= len(originals):
            sys.exit(f"plan slide {n}: 'from': {src} is out of range (deck has {len(originals)} slides)")
        result.append(originals[src - 1] if src not in used else None)
        used.add(src)
    # Copy first (from untouched sources), then fill: a copy must not inherit text of an earlier fill.
    for i, spec in enumerate(plan["slides"]):
        if result[i] is None:
            result[i] = duplicate_slide(prs, originals[int(spec["from"]) - 1])
    report = []
    for n, (slide, spec) in enumerate(zip(result, plan["slides"]), 1):
        try:
            apply(slide, spec, report)
        except (ValueError, OSError) as exc:
            sys.exit(f"plan slide {n} (from {spec['from']}): {exc} — nothing was saved.")
    set_order(prs, result)
    prs.save(out_path)
    print(f"saved {out_path}: {len(result)} slides "
          f"(from source slides {', '.join(str(s['from']) for s in plan['slides'])})")
    for line in report:
        print(line)
    check = os.path.join(os.path.dirname(os.path.abspath(__file__)), "pptx_inventory.py")
    if os.path.exists(check):
        sys.stdout.flush()
        sys.exit(subprocess.call([sys.executable, check, out_path, "--check"]))


if __name__ == "__main__":
    main()
